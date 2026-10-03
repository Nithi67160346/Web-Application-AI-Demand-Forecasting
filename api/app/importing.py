"""Bounded CSV/XLSX parsing and full-file quality reports before any sales mutation."""
import csv
import io
import zipfile
from datetime import date, datetime
from pydantic import ValidationError
from .sales import SaleInput

FIELDS = ['sale_date', 'product_code', 'region', 'sales_quantity', 'inventory']
MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 250000


def read_file(content, filename):
    if len(content) > MAX_BYTES:
        raise ValueError('File exceeds 20 MB.')
    if filename.lower().endswith('.csv'):
        reader = csv.reader(io.StringIO(content.decode('utf-8-sig')))
        header = next(reader, [])
        source = reader
    elif filename.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Expanded Excel archive exceeds 100 MB.')
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        source = workbook.worksheets[0].iter_rows(values_only=True)
        header = list(next(source, []))
    else:
        raise ValueError('Use UTF-8 CSV or Excel .xlsx (first sheet). Legacy .xls is unsupported.')
    try:
        header = [str(item or '').strip() for item in header]
        if not header or len(header) > 100 or any(not key for key in header) or len(set(header)) != len(header):
            raise ValueError('Headers must be nonempty, unique, and at most 100 columns.')
        rows = []
        for row in source:
            if not any(value is not None and str(value).strip() for value in row):
                continue
            if len(rows) >= MAX_ROWS:
                raise ValueError('At most 250,000 rows per import.')
            if len(row) > len(header) and any(str(value or '').strip() for value in row[len(header):]):
                raise ValueError('A row contains more columns than the header.')
            rows.append({key: (value.isoformat() if isinstance(value, (datetime, date)) else str(value or '') if value != 0 else '0')
                         for key, value in zip(header, list(row) + [''] * max(0, len(header)-len(row)))})
        if not rows:
            raise ValueError('File contains no data rows.')
        return header, rows
    finally:
        if filename.lower().endswith('.xlsx'):
            workbook.close()


def validate_rows(header, rows, mapping):
    for field in FIELDS[:4]:
        if mapping.get(field) not in header:
            raise ValueError(f'Map required column {field}.')
    mapped = [value for value in mapping.values() if value]
    if len(set(mapped)) != len(mapped) or any(value not in header for value in mapped):
        raise ValueError('Each source column can be mapped once and must exist in the file.')
    valid, errors, seen = [], [], set()
    invalid = duplicates = missing_inventory = 0
    for index, raw in enumerate(rows, 2):
        item = {field: raw.get(mapping.get(field, ''), '') for field in FIELDS}
        try:
            item['sale_date'] = str(item['sale_date']).split('T')[0]
            for field in ('sales_quantity', 'inventory'):
                value = str(item[field]).strip()
                if field == 'inventory' and not value:
                    item[field] = None
                    missing_inventory += 1
                else:
                    number = float(value)
                    if not number.is_integer():
                        raise ValueError(f'{field} must be an integer.')
                    item[field] = int(number)
            normalized = SaleInput.model_validate(item).model_dump(mode='json')
            key = tuple(normalized[field] for field in FIELDS)
            if key in seen:
                duplicates += 1
            seen.add(key)
            valid.append(normalized)
        except (ValidationError, ValueError, OverflowError) as exc:
            invalid += 1
            if len(errors) < 100:
                errors.append({'row': index, 'error': str(exc)[:400]})
    dates = [row['sale_date'] for row in valid]
    report = {'total_rows': len(rows), 'valid_rows': len(valid), 'invalid_rows': invalid,
              'exact_duplicate_rows': duplicates, 'missing_inventory': missing_inventory,
              'start_date': min(dates) if dates else None, 'end_date': max(dates) if dates else None,
              'errors': errors, 'quality_percent': round(100 * len(valid) / len(rows), 2),
              'duplicate_policy': 'Repeated transactions can be legitimate. No automatic deduplication; explicitly acknowledge duplicates before importing.',
              'inventory_policy': 'Inventory is a dated stock snapshot. Never sum repeated daily transaction snapshots.',
              'preview': valid[:20], 'headers': header, 'mapping': mapping}
    return valid, report
